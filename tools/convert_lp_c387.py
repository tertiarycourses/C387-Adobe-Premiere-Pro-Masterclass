"""C387 LP: WSQ TGS-2026064536 -> non-WSQ C387, 2 days 9:30am-5:30pm, 7.5 h/day, no assessment."""
import sys; sys.path.insert(0, 'tools')
from docx import Document
from docx_common import *

SRC = 'source-wsq/courseware/LP-Video Editing with Premiere Pro.docx'
OUT = 'courseware/LP-Adobe Premiere Pro Masterclass.docx'
d = Document(SRC)
P = d.paragraphs
def para(prefix):
    return next(p for p in P if p.text.startswith(prefix))

# cover + TOC
set_para(para('TGS Ref No:'), 'Course Code: C387')
toc = para('Two-Day Schedule  ·')
toc.runs[1].text = '  ·  15 instructional hours (9:30am – 5:30pm daily)'
for r in toc.runs[2:]: r.text = ''
# version record row
set_cell(d.tables[0].rows[2].cells[2], 'Non-WSQ edition: 15-hour two-day plan (9:30am – 5:30pm) for five topics and ten Premiere labs.')
# course information
info = d.tables[1]
for r in info.rows:
    k = r.cells[0].text
    if k == 'Duration': set_cell(r.cells[1], '2 days · 15 instructional hours (9:30am – 5:30pm daily)')
    if k == 'Delivery': set_cell(r.cells[1], '15 hours facilitated learning and hands-on Premiere Pro labs')
tsc = next(r for r in info.rows if r.cells[0].text == 'TSC')
delete_el(tsc._tr)
# Assessment section (heading + bullets)
h = para('Assessment')
assert h.style.name.startswith('Heading')
nxt = h._p.getnext()
delete_el(h._p)
while nxt is not None and ''.join(t.text or '' for t in nxt.iter(qn('w:t'))) != 'Two-Day Schedule':
    n2 = nxt.getnext(); delete_el(nxt); nxt = n2
# day headings
for p in P:
    if p.style.name == 'Heading 2' and p.text in ('Day 1', 'Day 2'):
        set_para(p, p.text + ' · 9:30am – 5:30pm')

# schedules
d1, d2 = d.tables[2], d.tables[3]
s1 = [[c.text for c in r.cells] for r in d1.rows[1:]]
s2 = [[c.text for c in r.cells] for r in d2.rows[1:]]
rebuild_rows(d1, [
    ('09:30–10:00', '30 min', 'Admin', 'Introductions, ground rules, course materials and lab setup', 'Slides 1–10'),
    ('10:00–11:05', '65 min', s1[1][2], s1[1][3], 'Topic 1: slide 12'),
    ('11:05–12:55', '110 min', s1[3][2], s1[3][3], 'Labs 1–2: 32, 33'),
    ('12:55–13:25', '30 min', 'Lunch', 'Lunch break', '—'),
    ('13:25–15:10', '105 min', s1[5][2], s1[5][3], 'Topic 2: 35 · Labs: 57, 58'),
    ('15:10–17:00', '110 min', s1[7][2], s1[7][3], 'Topic 3: 60 · Lab 5: 82'),
    ('17:00–17:30', '30 min', 'Checkpoint', 'Day 1 review and editing-log audit', 'Topic checkpoints'),
])
rebuild_rows(d2, [
    ('09:30–09:55', '25 min', 'Admin', 'Day 1 recap', 'Recap'),
    ('09:55–11:45', '110 min', s2[1][2], s2[1][3], 'Labs 6–7: 83, 84'),
    ('11:45–12:30', '45 min', 'Topic 4', 'Speed, time remapping, scopes and colour correction', 'Topic 4: 86'),
    ('12:30–13:00', '30 min', 'Lunch', 'Lunch break', '—'),
    ('13:00–14:45', '105 min', 'Topic 4 + Labs', 'Colour matching, review and collaboration; Labs 8–9', 'Labs: 104, 105'),
    ('14:45–16:20', '95 min', s2[5][2], s2[5][3], 'Topic 5: 107 · Lab 10: 126'),
    ('16:20–17:30', '70 min', 'Synthesis', 'Capstone review, course synthesis and course feedback', 'Closing slides 128–131'),
])
# lab alignment: remap slide column
for t in d.tables[4:6]:
    for r in t.rows[1:]:
        set_cell(r.cells[2], str(remap(int(r.cells[2].text))))
global_subs(d)
d.save(OUT)
print('saved', OUT)
