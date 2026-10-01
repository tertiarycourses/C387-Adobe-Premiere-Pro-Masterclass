"""C387 LG: WSQ TGS-2026064536 -> non-WSQ C387. Drops TSC alignment + assessment sources; keeps LOs."""
import sys; sys.path.insert(0, 'tools')
from docx import Document
from docx_common import *

SRC = 'source-wsq/courseware/LG-Video Editing with Premiere Pro.docx'
OUT = 'courseware/LG-Adobe Premiere Pro Masterclass.docx'
d = Document(SRC)
P = d.paragraphs
def para(prefix):
    return next(p for p in P if p.text.startswith(prefix))

set_para(para('TGS Ref No:'), 'Course Code: C387')
toc = para('Learning Outcomes and TSC Alignment')
toc.runs[0].text = 'Learning Outcomes'
toc.runs[1].text = '  ·  LO1–LO5'
for r in toc.runs[2:]: r.text = ''
# TSC Alignment heading + K/A table
h = next(p for p in P if p.text == 'TSC Alignment' and p.style.name.startswith('Heading'))
tbl = h._p.getnext()
assert tbl.tag == qn('w:tbl')
delete_el(h._p); delete_el(tbl)
# Source register: drop WSQ deck + assessment papers, repoint course page
reg = next(t for t in d.tables if t.rows[0].cells[0].text == 'Authority')
for r in list(reg.rows):
    k = r.cells[0].text
    if k in ('Legacy course deck', 'Original WA from LMS-TMS', 'Original PP from LMS-TMS'):
        delete_el(r._tr)
    elif k == 'Official course page':
        set_cell(r.cells[1], 'https://www.tertiarycourses.com.sg/adobe-premiere-pro-masterclass.html')
global_subs(d)
d.save(OUT)
print('saved', OUT)
