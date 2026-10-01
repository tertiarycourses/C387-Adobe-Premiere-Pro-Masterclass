import copy, re
from docx.oxml.ns import qn

OLD_T, NEW_T = 'Video Editing with Premiere Pro', 'Adobe Premiere Pro Masterclass'
OLD_C, NEW_C = 'TGS-2026064536', 'C387'
DROPPED = [2, 9, 10, 11, 12, 14, 15, 16, 139, 140]   # 1-based source slides removed from the deck

def remap(n):
    assert n not in DROPPED, n
    return n - sum(1 for d in DROPPED if d < n)

def remap_text(t):
    return re.sub(r'\b(\d{2,3})\b', lambda m: str(remap(int(m.group(1)))) if 17 <= int(m.group(1)) <= 142 else m.group(0), t)

def set_para(pa, text):
    if not pa.runs:
        pa.add_run(text); return
    pa.runs[0].text = text
    for r in pa.runs[1:]:
        r.text = ''

def set_cell(cell, text):
    set_para(cell.paragraphs[0], text)

def strip_ka(t):
    t = re.sub(r'(LO\d)( · [KA]\d)+', r'\1', t)
    return t

def all_paras(doc):
    for pa in doc.paragraphs:
        yield pa
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                yield from c.paragraphs
    for s in doc.sections:
        for part in (s.header, s.footer, s.even_page_header, s.even_page_footer, s.first_page_header, s.first_page_footer):
            yield from part.paragraphs
            for t in part.tables:
                for r in t.rows:
                    for c in r.cells:
                        yield from c.paragraphs

def global_subs(doc):
    for pa in all_paras(doc):
        for r in pa.runs:
            n = strip_ka(r.text.replace(OLD_T, NEW_T).replace(OLD_C, NEW_C))
            if n != r.text:
                r.text = n
        full = pa.text
        n = strip_ka(full.replace(OLD_T, NEW_T).replace(OLD_C, NEW_C))
        if n != full:
            set_para(pa, n)

def delete_el(el):
    el.getparent().remove(el)

def rebuild_rows(table, rows):
    """Replace all body rows with `rows` (list of 5-tuples), cloning the first body row."""
    tpl = table.rows[1]._tr
    olds = [r._tr for r in table.rows[1:]]
    for vals in rows:
        tr = copy.deepcopy(tpl)
        olds[-1].addnext(tr) if False else table._tbl.append(tr)
    for o in olds:
        delete_el(o)
    for r, vals in zip(table.rows[1:], rows):
        for c, v in zip(r.cells, vals):
            set_cell(c, v)
